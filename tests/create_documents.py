"""Create valid office fixtures with synthetic content, never user documents."""
from pathlib import Path
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1] / 'test-results' / 'documents'
ROOT.mkdir(parents=True, exist_ok=True)
document = Document()
document.add_heading('会议纪要测试', 0)
document.add_paragraph('普通办公文档 · 中文、标点与 Unicode 😀')
document.save(ROOT / '会议纪要.docx')
book = Workbook()
sheet = book.active
sheet.title = '统计'
sheet.append(['项目', '金额'])
sheet.append(['办公用品', 128.5])
sheet.append(['合计', '=SUM(B2:B2)'])
book.save(ROOT / '统计表.xlsx')
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text = 'PPTX 测试文稿'
slide.placeholders[1].text = '文件内容和格式保持原样'
slides.save(ROOT / '演示文稿.pptx')
Image.new('RGBA', (3, 2), (30, 100, 220, 128)).save(ROOT / '示意图.png')
(ROOT / '普通文本.txt').write_bytes('会议记录\r\n保留换行、中文与表情 😀\n'.encode('utf-8'))
(ROOT / '原始数据.bin').write_bytes(bytes(range(256)) * 31)
(ROOT / 'empty.txt').write_bytes(b'')

# A small valid PDF with correct object offsets and cross-reference table.
stream = b'BT /F1 12 Tf 40 90 Td (Synthetic office attachment) Tj ET'
objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
           b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
           b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 140] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
           b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
           b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
pdf = bytearray(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n')
offsets = [0]
for index, body in enumerate(objects, 1):
    offsets.append(len(pdf))
    pdf.extend(f'{index} 0 obj\n'.encode() + body + b'\nendobj\n')
start = len(pdf)
pdf.extend(b'xref\n0 6\n0000000000 65535 f \n')
for offset in offsets[1:]:
    pdf.extend(f'{offset:010d} 00000 n \n'.encode())
pdf.extend(f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n'.encode())
(ROOT / '附件.pdf').write_bytes(pdf)
print('Created valid DOCX/XLSX/PPTX/PDF/PNG/text/binary/empty fixtures.')
