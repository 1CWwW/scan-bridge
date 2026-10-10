"""Verify received bytes independently and reopen the generated Office files."""
from pathlib import Path
import hashlib
import json
import zipfile
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from PIL import Image

root = Path(__file__).resolve().parents[1] / 'test-results' / 'documents'
files = ['会议纪要.docx', '统计表.xlsx', '演示文稿.pptx', '附件.pdf', '示意图.png', '普通文本.txt', '原始数据.bin']
for name in files:
    assert (root / name).read_bytes() == (root / ('received-' + name)).read_bytes(), name
with zipfile.ZipFile(root / 'office-package.zip') as archive:
    assert archive.testzip() is None
    for name in files:
        assert archive.read(name) == (root / name).read_bytes(), name
    assert archive.read('办公资料/build/附件.bin') == bytes([0, 255, 128, 10])
    assert archive.read('办公资料/empty.txt') == b''
    assert '__scanbridge__/text-summary.md' not in archive.namelist()
    manifest = json.loads(archive.read('__scanbridge__/manifest.json'))
    assert len(manifest['files']) == 9
    for file in manifest['files']:
        raw = archive.read(file['path'])
        assert file['bytes'] == len(raw)
        assert file['sha256'] == hashlib.sha256(raw).hexdigest()
assert Document(root / 'received-会议纪要.docx').paragraphs[0].text == '会议纪要测试'
assert load_workbook(root / 'received-统计表.xlsx')['统计']['B3'].value == '=SUM(B2:B2)'
assert Presentation(root / 'received-演示文稿.pptx').slides[0].shapes.title.text == 'PPTX 测试文稿'
with Image.open(root / 'received-示意图.png') as picture:
    assert picture.size == (3, 2) and picture.mode == 'RGBA'
print('PASS: original downloads and ZIP bytes/hashes; independent DOCX/XLSX/PPTX/image readers reopened received fixtures.')
